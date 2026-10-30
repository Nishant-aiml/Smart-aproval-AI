"""
Master Assembly Script for Smart Approval AI B.Tech Final Thesis
Priyadarshini College of Engineering, Nagpur
Department of Information Technology (Session 2026-2027)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Add local path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from thesis_prelim import build_preliminary_pages, add_paragraph
from thesis_ch1 import build_chapter_1
from thesis_ch2 import build_chapter_2
from thesis_ch3 import build_chapter_3
from thesis_ch4 import build_chapter_4
from thesis_ch5 import build_chapter_5
from thesis_ch6 import build_chapter_6
from thesis_endmatter import build_end_matter
from thesis_additions import (
    add_chapter_2_deep_dive,
    add_chapter_3_rubrics_deep_dive,
    add_chapter_4_ddl_schemas,
    add_chapter_5_ablation_studies
)
from thesis_additions_part2 import (
    add_chapter_3_naac_and_trace,
    add_chapter_4_middleware_and_proxy,
    add_chapter_5_case_studies_and_roi,
    add_appendix_code_listings
)
from thesis_additions_part3 import (
    add_chapter_3_nba_code,
    add_chapter_5_security_and_pen_testing,
    add_chapter_6_concluding_remarks
)
from thesis_additions_part4 import (
    add_chapter_2_protocols,
    add_chapter_4_docker,
    add_chapter_5_sus
)

OUTPUT_DOCX = "thesis/Smart_Approval_AI_Final_Thesis.docx"

def apply_page_border(sect):
    border_xml = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'<w:top w:val="single" w:color="000000" w:sz="18" w:space="24"/>'
        f'<w:left w:val="single" w:color="000000" w:sz="18" w:space="24"/>'
        f'<w:bottom w:val="single" w:color="000000" w:sz="18" w:space="24"/>'
        f'<w:right w:val="single" w:color="000000" w:sz="18" w:space="24"/>'
        f'</w:pgBorders>'
    )
    sect._sectPr.append(border_xml)

def set_section_margins(sect):
    sect.top_margin = Inches(1.0)
    sect.bottom_margin = Inches(1.0)
    sect.left_margin = Inches(1.5)
    sect.right_margin = Inches(1.0)

def add_dynamic_page_number(run):
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def generate_thesis():
    print("=" * 60)
    print("STARTING FULL THESIS GENERATION...")
    print("=" * 60)
    
    doc = docx.Document()
    
    # Configure base style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    
    # Section 0: Preliminary Pages
    sect_prelim = doc.sections[0]
    set_section_margins(sect_prelim)
    apply_page_border(sect_prelim)
    
    # Configure footer for preliminary pages
    f_prelim = sect_prelim.footer
    f_prelim.is_linked_to_previous = False
    p_fp = f_prelim.paragraphs[0]
    p_fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fp.text = ""
    r_fp = p_fp.add_run()
    r_fp.font.name = "Times New Roman"
    r_fp.font.size = Pt(11)
    # Front matter dynamic page number
    add_dynamic_page_number(r_fp)
    
    # Configure header
    h_prelim = sect_prelim.header
    h_prelim.is_linked_to_previous = False
    p_hp = h_prelim.paragraphs[0]
    p_hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hp = p_hp.add_run("Priyadarshini College of Engineering, Nagpur | Department of Information Technology")
    r_hp.font.name = "Times New Roman"
    r_hp.font.size = Pt(9)
    r_hp.font.italic = True
    r_hp.font.color.rgb = RGBColor(120, 120, 120)
    
    print("Building Preliminary Pages...")
    build_preliminary_pages(doc)
    
    # Section 1: Main Content Chapters
    sect_main = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    set_section_margins(sect_main)
    apply_page_border(sect_main)
    
    f_main = sect_main.footer
    f_main.is_linked_to_previous = False
    p_fm = f_main.paragraphs[0]
    p_fm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fm.text = ""
    r_fm = p_fm.add_run()
    r_fm.font.name = "Times New Roman"
    r_fm.font.size = Pt(11)
    add_dynamic_page_number(r_fm)
    
    h_main = sect_main.header
    h_main.is_linked_to_previous = False
    p_hm = h_main.paragraphs[0]
    p_hm.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hm = p_hm.add_run("Smart Approval AI: B.Tech Project Dissertation (2026-2027)")
    r_hm.font.name = "Times New Roman"
    r_hm.font.size = Pt(9)
    r_hm.font.italic = True
    r_hm.font.color.rgb = RGBColor(120, 120, 120)
    
    print("Building Chapter 1: Introduction...")
    build_chapter_1(doc)
    
    print("Building Chapter 2: Literature Review...")
    build_chapter_2(doc)
    add_chapter_2_deep_dive(doc)
    add_chapter_2_protocols(doc)
    
    print("Building Chapter 3: Methodology & Proposed System Architecture...")
    build_chapter_3(doc)
    add_chapter_3_rubrics_deep_dive(doc)
    add_chapter_3_naac_and_trace(doc)
    add_chapter_3_nba_code(doc)
    
    print("Building Chapter 4: Hardware & Software Requirements...")
    build_chapter_4(doc)
    add_chapter_4_ddl_schemas(doc)
    add_chapter_4_middleware_and_proxy(doc)
    add_chapter_4_docker(doc)
    
    print("Building Chapter 5: Results & Discussions...")
    build_chapter_5(doc)
    add_chapter_5_ablation_studies(doc)
    add_chapter_5_case_studies_and_roi(doc)
    add_chapter_5_security_and_pen_testing(doc)
    add_chapter_5_sus(doc)
    
    print("Building Chapter 6: Conclusion & Future Scope...")
    build_chapter_6(doc)
    add_chapter_6_concluding_remarks(doc)
    
    print("Building References & Appendices...")
    build_end_matter(doc)
    add_appendix_code_listings(doc)
    
    # Ensure all sections have borders applied
    for idx, s in enumerate(doc.sections):
        set_section_margins(s)
        apply_page_border(s)
        
    print(f"Saving compiled thesis to: {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    
    file_size_mb = os.path.getsize(OUTPUT_DOCX) / (1024 * 1024)
    print("=" * 60)
    print("THESIS GENERATION COMPLETE!")
    print(f"File Saved: {OUTPUT_DOCX}")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print(f"Total Sections: {len(doc.sections)}")
    print(f"Output File Size: {file_size_mb:.2f} MB")
    print("=" * 60)

if __name__ == "__main__":
    generate_thesis()
